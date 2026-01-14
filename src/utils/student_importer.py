"""
Excel'den Öğrenci Listesi İçe Aktarma Sınıfı
"""

import os
from typing import List, Dict, Tuple, Optional, Set
from dataclasses import dataclass
from datetime import datetime

import openpyxl
import xlrd

from src.models.student import Student, StudentCourse
from src.repositories.student_repository import StudentRepository, StudentCourseRepository
from src.repositories.course_repository import CourseRepository
from src.repositories.department_repository import DepartmentRepository


@dataclass
class ImportResult:
    """İçe aktarma sonucu"""
    success: bool
    message: str
    students_imported: int = 0
    student_courses_created: int = 0
    errors: List[str] = None
    warnings: List[str] = None
    
    def __post_init__(self):
        if self.errors is None:
            self.errors = []
        if self.warnings is None:
            self.warnings = []


class StudentImporter:
    
    TURKISH_CHAR_MAP = {
        'ı': 'i', 'İ': 'I',
        'ğ': 'g', 'Ğ': 'G',
        'ü': 'u', 'Ü': 'U',
        'ş': 's', 'Ş': 'S',
        'ö': 'o', 'Ö': 'O',
        'ç': 'c', 'Ç': 'C'
    }
    
    DEPARTMENT_CODE_MAP = {
        'BLM': 'BLM',      # Bilgisayar Mühendisliği
        'BILM': 'BLM',
        'MAT': 'MAT',      # Matematik Mühendisliği
        'YZM': 'YZM',      # Yazılım Mühendisliği (eğer varsa)
        'SEC': 'SEC',      # SecMühendisliği (eğer varsa)
        'EE': 'EE',        # Elektrik-Elektronik
        'ELK': 'ELK',
        'ME': 'ME',        # Makine Mühendisliği
        'MAK': 'MAK',
        'CE': 'CE',        # İnşaat Mühendisliği
        'INS': 'INS',
    }
    
    def __init__(self):
        self.student_repo = StudentRepository()
        self.student_course_repo = StudentCourseRepository()
        self.course_repo = CourseRepository()
        self.department_repo = DepartmentRepository()
    
    def import_from_excel(
        self,
        file_path: str,
        course_id: Optional[int] = None,
        course_code: Optional[str] = None,
        department_id: Optional[int] = None,
        semester: Optional[str] = None,
        year: Optional[int] = None
    ) -> ImportResult:
        if not os.path.exists(file_path):
            return ImportResult(
                success=False,
                message=f"Dosya bulunamadı: {file_path}"
            )
        
        # Course ID'yi belirle
        target_course_id = course_id
        if not target_course_id and course_code:
            course = self.course_repo.get_by_code(course_code)
            if course:
                target_course_id = course.id
                department_id = department_id or course.department_id
                year = year or course.year
        
        if not target_course_id:
            return ImportResult(
                success=False,
                message="Ders belirtilmedi veya bulunamadı"
            )
        
        try:
            students_data = self._read_excel_file(file_path, department_id, year)
        except Exception as e:
            return ImportResult(
                success=False,
                message=f"Excel dosyası okunamadı: {str(e)}"
            )
        
        if not students_data:
            return ImportResult(
                success=False,
                message="Öğrenci verisi bulunamadı"
            )
        
        result = self._import_students(students_data, target_course_id, semester)
        
        return result
    
    def import_from_excel_directory(
        self,
        directory_path: str,
        semester: Optional[str] = None,
        department_id: Optional[int] = None
    ) -> Dict[str, ImportResult]:
        results = {}
        
        if not os.path.exists(directory_path):
            return {'': ImportResult(False, "Klasör bulunamadı")}
        
        for filename in os.listdir(directory_path):
            if not (filename.endswith('.xls') or filename.endswith('.xlsx')):
                continue
            
            if filename.endswith('Zone.Identifier'):
                continue
            
            file_path = os.path.join(directory_path, filename)
            
            course_code = self._extract_course_code_from_filename(filename)
            
            result = self.import_from_excel(
                file_path=file_path,
                course_code=course_code,
                department_id=department_id,
                semester=semester
            )
            
            results[filename] = result
        
        return results
    
    def _read_excel_file(
        self,
        file_path: str,
        department_id: Optional[int] = None,
        year: Optional[int] = None
    ) -> List[Dict]:

        students = []
        
        if file_path.endswith('.xlsx'):
            return self._read_xlsx(file_path, department_id, year)
        else:  # .xls
            return self._read_xls(file_path, department_id, year)
    
    def _read_xls(
        self,
        file_path: str,
        department_id: Optional[int] = None,
        year: Optional[int] = None
    ) -> List[Dict]:

        students = []
        
        try:
            workbook = xlrd.open_workbook(file_path, formatting_info=False)
            sheet = workbook.sheet_by_index(0)
            
            header_row = self._find_header_row_xls(sheet)
            if header_row is None:
                header_row = 0
            
            col_indices = self._get_column_indices_xls(sheet, header_row)
            
            for row_idx in range(header_row + 1, sheet.nrows):
                row_data = {}
                
                for key, col_idx in col_indices.items():
                    if col_idx is not None and col_idx < sheet.ncols:
                        cell_value = sheet.cell_value(row_idx, col_idx)
                        row_data[key] = self._clean_cell_value(cell_value)
                
                student_number = row_data.get('student_number', '').strip()
                if not student_number or not self._is_valid_student_number(student_number):
                    continue
                
                first_name = row_data.get('first_name', '').strip()
                last_name = row_data.get('last_name', '').strip()
                
                if not first_name and not last_name:
                    full_name = row_data.get('full_name', '').strip()
                    if full_name:
                        name_parts = full_name.split()
                        if len(name_parts) >= 2:
                            first_name = ' '.join(name_parts[:-1])
                            last_name = name_parts[-1]
                        else:
                            first_name = full_name
                
                if not first_name or not last_name:
                    continue
                
                student = {
                    'student_number': student_number,
                    'first_name': first_name,
                    'last_name': last_name,
                    'email': row_data.get('email', ''),
                    'department_id': department_id,
                    'year': year or self._parse_year(row_data.get('year', ''))
                }
                
                students.append(student)
                
        except Exception as e:
            raise Exception(f"XLS okuma hatası: {str(e)}")
        
        return students
    
    def _read_xlsx(
        self,
        file_path: str,
        department_id: Optional[int] = None,
        year: Optional[int] = None
    ) -> List[Dict]:
        students = []
        
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet = workbook.active
            
            header_row = self._find_header_row_xlsx(sheet)
            if header_row is None:
                header_row = 1
            
            col_indices = self._get_column_indices_xlsx(sheet, header_row)
            
            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                row_data = {}
                
                for key, col_idx in col_indices.items():
                    if col_idx is not None and col_idx < len(row):
                        row_data[key] = self._clean_cell_value(row[col_idx])
                
                student_number = row_data.get('student_number', '').strip()
                if not student_number or not self._is_valid_student_number(student_number):
                    continue
                
                first_name = row_data.get('first_name', '').strip()
                last_name = row_data.get('last_name', '').strip()
                
                if not first_name and not last_name:
                    full_name = row_data.get('full_name', '').strip()
                    if full_name:
                        name_parts = full_name.split()
                        if len(name_parts) >= 2:
                            first_name = ' '.join(name_parts[:-1])
                            last_name = name_parts[-1]
                        else:
                            first_name = full_name
                
                if not first_name or not last_name:
                    continue
                
                student = {
                    'student_number': student_number,
                    'first_name': first_name,
                    'last_name': last_name,
                    'email': row_data.get('email', ''),
                    'department_id': department_id,
                    'year': year or self._parse_year(row_data.get('year', ''))
                }
                
                students.append(student)
                
        except Exception as e:
            raise Exception(f"XLSX okuma hatası: {str(e)}")
        
        return students
    
    def _find_header_row_xls(self, sheet) -> Optional[int]:
        common_headers = ['ogrenci no', 'öğrenci no', 'numara', 'no', 'no.', 'sıra no']
        
        for row_idx in range(min(10, sheet.nrows)):
            for col_idx in range(sheet.ncols):
                cell_value = str(sheet.cell_value(row_idx, col_idx)).lower().strip()
                for header in common_headers:
                    if header in cell_value:
                        return row_idx
        
        return None
    
    def _find_header_row_xlsx(self, sheet) -> Optional[int]:
        common_headers = ['ogrenci no', 'öğrenci no', 'numara', 'no', 'no.', 'sıra no']
        
        for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True)):
            for cell_value in row:
                if cell_value:
                    cell_str = str(cell_value).lower().strip()
                    for header in common_headers:
                        if header in cell_str:
                            return row_idx + 1
        
        return None
    
    def _get_column_indices_xls(self, sheet, header_row: int) -> Dict[str, Optional[int]]:
        indices = {
            'student_number': None,
            'first_name': None,
            'last_name': None,
            'full_name': None,
            'email': None,
            'year': None
        }
        
        if header_row >= sheet.nrows:
            return indices
        
        for col_idx in range(sheet.ncols):
            header = str(sheet.cell_value(header_row, col_idx)).lower().strip()
            
            if any(x in header for x in ['öğrenci no', 'ogrenci no', 'numara', 'no.', 'no ']):
                indices['student_number'] = col_idx
            elif 'ad' in header and 'soyad' not in header:
                indices['first_name'] = col_idx
            elif 'soyad' in header:
                indices['last_name'] = col_idx
            elif 'ad soyad' in header or 'adı soyadı' in header:
                indices['full_name'] = col_idx
            elif 'e-posta' in header or 'eposta' in header or 'email' in header or 'e-post' in header:
                indices['email'] = col_idx
            elif 'sınıf' in header:
                indices['year'] = col_idx
        
        return indices
    
    def _get_column_indices_xlsx(self, sheet, header_row: int) -> Dict[str, Optional[int]]:
        indices = {
            'student_number': None,
            'first_name': None,
            'last_name': None,
            'full_name': None,
            'email': None,
            'year': None
        }
        
        header_values = list(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))[0]
        
        for col_idx, header in enumerate(header_values):
            if header:
                header_str = str(header).lower().strip()
                
                if any(x in header_str for x in ['öğrenci no', 'ogrenci no', 'numara', 'no.', 'no ']):
                    indices['student_number'] = col_idx
                elif 'ad' in header_str and 'soyad' not in header_str and 'ad soyad' not in header_str:
                    indices['first_name'] = col_idx
                elif 'soyad' in header_str:
                    indices['last_name'] = col_idx
                elif 'ad soyad' in header_str or 'adı soyadı' in header_str:
                    indices['full_name'] = col_idx
                elif any(x in header_str for x in ['e-posta', 'eposta', 'email', 'e-post']):
                    indices['email'] = col_idx
                elif 'sınıf' in header_str:
                    indices['year'] = col_idx
        
        return indices
    
    def _clean_cell_value(self, value) -> str:
        if value is None:
            return ''
        
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        
        return str(value).strip()
    
    def _is_valid_student_number(self, number: str) -> bool:
        if not number:
            return False
        
        number = number.strip()
        if not number.isdigit():
            return False
        
        return len(number) >= 3
    
    def _parse_year(self, year_str: str) -> int:
        if not year_str:
            return 1
        
        year_str = str(year_str).lower().strip()
        
        roman_map = {'i': 1, 'ii': 2, 'iii': 3, 'iv': 4, 'v': 5, 'vi': 6}
        if year_str in roman_map:
            return roman_map[year_str]
        
        if year_str.isdigit():
            year_int = int(year_str)
            return 1 if year_int < 1 or year_int > 6 else year_int
        
        if '1' in year_str or 'bir' in year_str:
            return 1
        elif '2' in year_str or 'iki' in year_str:
            return 2
        elif '3' in year_str or 'üç' in year_str:
            return 3
        elif '4' in year_str or 'dört' in year_str:
            return 4
        
        return 1
    
    def _extract_course_code_from_filename(self, filename: str) -> Optional[str]:
        import re
        
        match = re.search(r'\[([A-Z0-9]+)\]', filename)
        if match:
            return match.group(1)
        
        match = re.search(r'([A-Z]{3,}\d{3})', filename)
        if match:
            return match.group(1)
        
        return None
    
    def _import_students(
        self,
        students_data: List[Dict],
        course_id: int,
        semester: Optional[str] = None
    ) -> ImportResult:
        errors = []
        warnings = []
        
        students_to_create = []
        student_courses_to_create = []
        
        for data in students_data:
            try:
                student = Student(
                    student_number=data['student_number'],
                    first_name=data['first_name'],
                    last_name=data['last_name'],
                    email=data.get('email') or None,
                    department_id=data.get('department_id'),
                    year=data.get('year', 1),
                    is_active=True
                )
                
                students_to_create.append(student)
                
            except Exception as e:
                errors.append(f"{data.get('student_number', 'Bilinmeyen')}: {str(e)}")
        
        try:
            created_count = self.student_repo.create_batch(students_to_create)
            
            imported_students = []
            for data in students_data:
                student = self.student_repo.get_by_student_number(data['student_number'])
                if student:
                    imported_students.append(student)
            
            for student in imported_students:
                student_course = StudentCourse(
                    student_id=student.id,
                    course_id=course_id,
                    semester=semester,
                    is_active=True
                )
                student_courses_to_create.append(student_course)
            
            if student_courses_to_create:
                self.student_course_repo.create_batch(student_courses_to_create)
            
            return ImportResult(
                success=True,
                message=f"{len(imported_students)} öğrenci başarıyla içe aktarıldı",
                students_imported=len(imported_students),
                student_courses_created=len(student_courses_to_create),
                errors=errors if errors else None,
                warnings=warnings if warnings else None
            )
            
        except Exception as e:
            return ImportResult(
                success=False,
                message=f"Kayıt sırasında hata: {str(e)}",
                errors=errors
            )
    
    def get_students_by_course(self, course_id: int) -> List[Student]:
        student_courses = self.student_course_repo.get_by_course_id(course_id)
        students = []
        
        for sc in student_courses:
            if sc.student_id:
                student = self.student_repo.get_by_id(sc.student_id)
                if student:
                    students.append(student)
        
        return students
    
    def get_student_course_conflicts(self, course_id1: int, course_id2: int) -> int:
        return self.student_course_repo.check_student_overlap(course_id1, course_id2)
    
    def get_all_conflicts_for_course(self, course_id: int, min_overlap: int = 1) -> List[Dict]:
        return self.student_course_repo.get_conflicting_courses(course_id, min_overlap)
